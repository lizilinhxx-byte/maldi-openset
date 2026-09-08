from __future__ import annotations

import re
import unicodedata
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .util import stable_id


def normalize_label(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"[_/\\]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _species_table(workbook) -> list[tuple[str, int]]:
    sheet = workbook["strains per X"]
    rows: list[tuple[str, int]] = []
    for row in sheet.iter_rows(min_row=6, values_only=True):
        species = normalize_label(row[3] if len(row) > 3 else "")
        count = row[4] if len(row) > 4 else None
        if species and isinstance(count, (int, float)):
            rows.append((species, int(count)))
    if not rows:
        raise ValueError("no species rows found in taxonomy workbook")
    return rows


def match_species(strain_name: str, species_names: list[str]) -> str | None:
    normalized = normalize_label(strain_name).casefold()
    candidates = [
        species
        for species in species_names
        if normalized == species.casefold()
        or normalized.startswith(species.casefold() + " ")
    ]
    return max(candidates, key=len) if candidates else None


def load_taxonomy_workbook(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return spectrum-level and species-level metadata from the RKI workbook."""
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    species_rows = _species_table(workbook)
    species_names = [name for name, _ in species_rows]
    species_df = pd.DataFrame(species_rows, columns=["species", "declared_strain_count"])
    species_df["genus"] = species_df["species"].str.split().str[0]

    spectrum_rows: list[dict] = []
    sheet = workbook["list of spectra"]
    per_strain_ordinal: Counter[str] = Counter()
    for row in sheet.iter_rows(min_row=6, values_only=True):
        ordinal = row[0] if row else None
        strain_name = normalize_label(row[1] if len(row) > 1 else "")
        if not isinstance(ordinal, (int, float)) or not strain_name:
            continue
        species = match_species(strain_name, species_names)
        per_strain_ordinal[strain_name] += 1
        spectrum_rows.append(
            {
                "spectrum_ordinal": int(ordinal),
                "spectrum_id": f"rki_{int(ordinal):06d}",
                "strain_name": strain_name,
                "strain_id": stable_id(strain_name, prefix="strain_"),
                "replicate_ordinal": per_strain_ordinal[strain_name],
                "species": species,
                "genus": species.split()[0] if species else None,
                "taxonomy_status": "matched" if species else "unresolved",
            }
        )
    spectra = pd.DataFrame(spectrum_rows).sort_values("spectrum_ordinal").reset_index(drop=True)
    if spectra.empty:
        raise ValueError("no spectrum rows found in taxonomy workbook")
    species_df["observed_strain_count"] = species_df["species"].map(
        spectra.dropna(subset=["species"])
        .drop_duplicates(["species", "strain_id"])
        .groupby("species")["strain_id"]
        .size()
    ).fillna(0).astype(int)
    species_df["observed_spectrum_count"] = species_df["species"].map(
        spectra.groupby("species")["spectrum_id"].size()
    ).fillna(0).astype(int)
    return spectra, species_df


def assign_analysis_sets(
    manifest: pd.DataFrame,
    known_min_strains: int = 5,
    sensitivity_min_strains: int = 3,
) -> pd.DataFrame:
    required = {"species", "genus", "strain_id"}
    missing = required.difference(manifest.columns)
    if missing:
        raise KeyError(f"manifest is missing columns: {sorted(missing)}")
    result = manifest.copy()
    counts = (
        result.dropna(subset=["species"])
        .drop_duplicates(["species", "strain_id"])
        .groupby("species")["strain_id"]
        .size()
    )
    result["species_strain_count"] = result["species"].map(counts).fillna(0).astype(int)
    # The primary and sensitivity schemes overlap by design: species with 3-4
    # strains are unknown in the primary >=5-strain analysis but known in the
    # >=3-strain sensitivity analysis. Keep explicit booleans rather than
    # forcing the schemes into one mutually exclusive label.
    result["primary_known"] = result["species_strain_count"] >= known_min_strains
    result["primary_ood"] = ~result["primary_known"]
    result["sensitivity_known"] = result["species_strain_count"] >= sensitivity_min_strains
    result["analysis_set"] = np.where(result["primary_known"], "primary_known", "ood")
    known_genera = set(result.loc[result["analysis_set"] == "primary_known", "genus"].dropna())
    result["ood_distance"] = None
    result["ood_singleton"] = False
    is_ood = result["primary_ood"]
    result.loc[is_ood & result["genus"].isin(known_genera), "ood_distance"] = "near"
    result.loc[is_ood & ~result["genus"].isin(known_genera), "ood_distance"] = "far"
    result.loc[is_ood & (result["species_strain_count"] == 1), "ood_singleton"] = True
    result["ood_singleton"] = result["ood_singleton"].astype(bool)
    return result


def taxonomy_summary(manifest: pd.DataFrame) -> dict[str, int]:
    return {
        "spectra": int(len(manifest)),
        "strains": int(manifest["strain_id"].nunique()),
        "species": int(manifest["species"].nunique(dropna=True)),
        "genera": int(manifest["genus"].nunique(dropna=True)),
        "unresolved_spectra": int(manifest["species"].isna().sum()),
        "primary_known_species": int(
            manifest.loc[manifest["analysis_set"] == "primary_known", "species"].nunique()
        ),
        "primary_known_strains": int(
            manifest.loc[manifest["analysis_set"] == "primary_known", "strain_id"].nunique()
        ),
        "sensitivity_known_species": int(manifest.loc[manifest["sensitivity_known"], "species"].nunique()),
        "ood_species": int(manifest.loc[manifest["analysis_set"] == "ood", "species"].nunique()),
        "singleton_ood_species": int(
            manifest.loc[manifest["ood_singleton"], "species"].nunique()
        ),
    }
